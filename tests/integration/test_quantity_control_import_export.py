"""用户核心流程：实际工作簿导入→确认口径→多分包/分栋→同快照导出。"""
from decimal import Decimal
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from jiadun.core.contracts import run_contract
from jiadun.core.engine import quantity_control as qc
from jiadun.core.engine import settlement_io
from jiadun.core.export import excel_export
from jiadun.core.models import project as pm


def test_file_import_to_building_controls_and_registered_exports(tmp_path):
    info = pm.create_project('合成完整流程', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        periods = []
        for key, direction, kind, quantities in [
            ('U', 'upward', 'contract', (1000, 800)),
            ('A', 'downward', 'settlement', (600, 200)),
            ('B', 'downward', 'settlement', (500, 300)),
            ('C', 'downward', 'settlement', (0, 100)),
        ]:
            wb = Workbook()
            wb.remove(wb.active)
            for building, quantity in enumerate(quantities, 1):
                ws = wb.create_sheet(f'第1期{building}号楼')
                ws.append(['清单编码', '清单名称', '项目特征', '单位', '工程量', '综合单价', '合价'])
                ws.append(['C1', '混凝土', 'C30', 'm3', quantity, 1, quantity])
                if building == 1 and key in ('U', 'A'):
                    ws.append(['C2', '钢材', 'Q235', 't' if key == 'U' else 'kg',
                               1 if key == 'U' else 0.5, 1, 1 if key == 'U' else 0.5])
            source = tmp_path / f'{key}第1期.xlsx'
            wb.save(source)
            report = settlement_io.import_settlement_file(
                conn, info.project_id, Path(info.workspace_path), source,
                direction=direction, separate_files=True,
                document_category='upward_contract_boq' if kind == 'contract' else 'downward_settlement',
            )
            assert report.status == 'ok', report.message
            periods.append(report.period_id)
            context_id = qc.attach_period_context(
                conn, info.project_id, report.period_id, contract_key=key,
                business_period_no=1, unit_name=f'合成单位{key}', doc_kind=kind,
                work_scope='材料供应', actor='user', reason='合成原件范围明确',
            )
            for row in conn.execute(
                'SELECT li.id,rs.sheet_name FROM line_items li JOIN raw_sheets rs ON rs.id=li.sheet_id '
                'WHERE li.period_id=?', (report.period_id,),
            ).fetchall():
                building = '1号楼' if '1号楼' in row['sheet_name'] else '2号楼'
                qc.set_line_context(conn, info.project_id, row['id'], building=building,
                                    actor='user', reason='合成工作表范围确认')
                qc.confirm_line_context(conn, info.project_id, row['id'], actor='user', reason='合成楼栋确认')
            qc.confirm_period_context(conn, info.project_id, context_id, actor='user', reason='合成合同和计量口径确认')
        assert len(set(periods)) == 4
        active = run_contract.ensure_run_contract(conn, info.project_id)
        ledger = qc.build_quantity_ledger(conn, info.project_id)
        assert ledger['run_signature'] == active.signature
        controls = {row['building']: row for row in ledger['quantity_controls'] if row['code'] == 'C1'}
        assert Decimal(controls['1号楼']['downstream_quantity']) == Decimal('1100')
        assert Decimal(controls['1号楼']['delta_vs_contract']) == Decimal('100')
        assert controls['1号楼']['status_vs_contract'] == 'FAIL'
        assert controls[None]['status_vs_contract'] == controls[None]['status'] == 'PASS'
        assert Decimal(controls[None]['downstream_quantity']) == Decimal('1700')
        assert Decimal(controls[None]['baselines']['upward_contract']['quantity']) == Decimal('1800')
        assert {row['contract_key']: Decimal(row['quantity']) for row in controls[None]['contributions']} == {
            'A': Decimal('800'), 'B': Decimal('800'), 'C': Decimal('100'),
        }
        output = excel_export.export_workbook(conn, info.project_id, Path(info.workspace_path) / 'exports')
        book = load_workbook(output)
        assert {'工程量控制台账', '分包数量贡献', '有效数量分项', '有效数量与来源', '核量待确认范围'} <= set(book.sheetnames)
        rows = list(book['有效数量与来源'].iter_rows(values_only=True))
        headers = rows[0]
        data = [dict(zip(headers, row, strict=True)) for row in rows[1:]]
        steel = next(row for row in data if row['合同标识'] == 'A' and row['编码'] == 'C2')
        assert steel['原数量'] == '0.5' and steel['原单位'] == 'kg'
        assert steel['标准行数量'] == '0.0005'
        assert steel['楼栋'] == '1号楼' and steel['业务期号'] == 1
        assert all(row['运行签名'] == active.signature for row in data)
        word = excel_export.export_management_summary_docx(conn, info.project_id, Path(info.workspace_path) / 'exports')
        assert word.exists()
        document = Document(word)
        table = next(table for table in document.tables if table.rows[0].cells[0].text == '对象/范围/单位')
        first = [cell.text for cell in table.rows[1].cells]
        assert first[:4] == ['混凝土 / 1号楼 / m3', '1000', '待确认', '1100']
        assert '超出' in first[4] and '100m3' in first[4]
        whole = next([cell.text for cell in row.cells] for row in table.rows[1:]
                     if row.cells[0].text == '混凝土 / 全项目 / m3')
        assert whole[1] == '1800' and whole[3] == '1700'
        registered = conn.execute('SELECT run_signature FROM export_runs WHERE project_id=?', (info.project_id,)).fetchall()
        assert len(registered) == 2 and all(row['run_signature'] == active.signature for row in registered)
    finally:
        conn.close()
