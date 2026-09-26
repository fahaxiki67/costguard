"""CG-01 回归：累计归组的单位与业务兼容性。

预审缺陷（2026-09-26）：aggregate.group_key_of 仅用 code/name 归组，累计层
不检查 line_items.unit。同编码"1 吨×1000=1000 元"与"1000 千克×1=1000 元"
被直接累计为 cum_qty=1001 并生成跨单位加权平均单价。

修复口径（不改变全局 item_key，历史映射保持稳定）：
- 单位无法证明兼容 → 组/期 status='incomparable'，cum_qty、wavg_price 置
  None（绝不输出跨单位相加值），金额（货币口径）仍如实累计；
- 同单位同口径正例、别名单位正例必须保持原行为；
- 全部未注明单位沿用历史口径；部分注明无法证明兼容 → 不可比；
- 同组规格（feature）不一致 → 待复核（incomplete）+ 警示，不静默确认合并；
- 不可比状态沿 summary（界面）、导出累计表、待核实清单、对上对下对比表传播。
"""
import hashlib
import json
from decimal import Decimal
from pathlib import Path

import pytest

from jiadun.core.engine import aggregate
from jiadun.core.engine.aggregate import (
    group_key_of,
    normalize_unit,
    units_incompatible,
)

D = Decimal


def _build_project(tmp_path, rows):
    """构造最小结算项目。

    rows: list[(period_no, direction, code, name, feature, unit, qty, price, amount)]
    同一期内多行用相邻同 period_no 条目表达。
    """
    from jiadun.core.models import project as pm

    info = pm.create_project("CG01-单位兼容", tmp_path / "ws")
    info, conn = pm.open_project(Path(info.workspace_path))
    with conn:
        sheet_no = 0
        seen_periods: dict[tuple, tuple] = {}
        ordered: list[tuple] = []
        for row in rows:
            period_key = (row[0], row[1])  # (period_no, direction)
            if period_key not in seen_periods:
                seen_periods[period_key] = row[:4]
                ordered.append(row)
        # 逐期建期次/文件/sheet，再逐行写 line_items
        for period_no, direction, _code, _name, _feat, _unit, _qty, _price, _amt in ordered:
            content = f"cg01-unit-{period_no}-{direction}".encode()
            stored = tmp_path / f"src-{period_no}-{direction}.xlsx"
            stored.write_bytes(content)
            digest = hashlib.sha256(content).hexdigest()
            period_id = conn.execute(
                """INSERT INTO settlement_periods(project_id, period_no, title, direction)
                   VALUES (?,?,?,?)""",
                (info.project_id, period_no, f"第{period_no}期", direction),
            ).lastrowid
            sheet_no += 1
            file_id = conn.execute(
                """INSERT INTO source_files(project_id, original_path, stored_path,
                   original_name, sha256, size_bytes, file_type, imported_at)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (info.project_id, f"/cg01-{period_no}-{direction}.xlsx", str(stored),
                 f"cg01-{period_no}-{direction}.xlsx", digest, stored.stat().st_size,
                 "xlsx", "2026"),
            ).lastrowid
            batch_id = conn.execute(
                """INSERT INTO parse_batches(file_id, parser, parsed_at, status)
                   VALUES (?,?,?,?)""", (file_id, "test", "2026", "ok"),
            ).lastrowid
            sheet_id = conn.execute(
                """INSERT INTO raw_sheets(batch_id, sheet_index, sheet_name, n_rows,
                   n_cols, period_id) VALUES (?,?,?,?,?,?)""",
                (batch_id, sheet_no, f"第{period_no}期明细", 2, 6, period_id),
            ).lastrowid
            col_map = {"code": 1, "name": 2, "quantity": 3, "unit_price": 4, "amount": 5}
            conn.execute(
                """INSERT INTO table_headers(sheet_id, header_row_lo, header_row_hi,
                   col_map_json, confidence, needs_review, data_row_start, data_row_end)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (sheet_id, 1, 1, json.dumps(col_map), 1.0, 0, 2, 2),
            )
            for row in rows:
                if (row[0], row[1]) != (period_no, direction):
                    continue
                _, _d, code, name, feature, unit, qty, price, amount = row
                conn.execute(
                    """INSERT INTO line_items(period_id, sheet_id, code, name, feature,
                       unit, quantity, unit_price, amount, flags_json)
                       VALUES (?,?,?,?,?,?,?,?,?,?)""",
                    (period_id, sheet_id, code, name, feature, unit, qty, price, amount,
                     json.dumps({"row": 2})),
                )
    return info, conn


def _row(period_no, direction, code, name, feature, unit, qty, price, amount):
    return (period_no, direction, code, name, feature, unit, qty, price, amount)


# ---------- 纯函数 ----------


def test_units_incompatible_rules():
    assert units_incompatible({"t", "kg"})
    assert units_incompatible({"t", ""})  # 有单位与未注明混存：无法证明兼容
    assert not units_incompatible({"t"})  # 同一（归一化）单位：兼容
    assert not units_incompatible({""})  # 全部未注明：历史口径保持
    assert not units_incompatible(set())


def test_normalize_unit_aliases_match_anomalies_layer():
    assert normalize_unit("吨") == normalize_unit("T")
    assert normalize_unit("千克") == normalize_unit("kg")
    assert normalize_unit(" 千克 ") == "kg"


def test_group_key_unchanged():
    """归组键口径必须保持稳定：历史 period_totals.item_key 映射不失联。"""
    assert group_key_of("C1", "清单A") == "code:C1"
    assert group_key_of("", "清单A") == "name:清单A"
    assert group_key_of(None, "清单A") == "name:清单A"


# ---------- 核心缺陷：跨期不同单位 ----------


def test_same_code_different_units_never_summed(tmp_path):
    """预审主案例：1吨 + 1000千克 ≠ 1001，加权平均单价不可定义。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "C1", "清单A", "", "千克", "1000", "1", "1000"),
    ])
    try:
        aggs = aggregate.aggregate_project(conn, info.project_id, direction="downward")
        assert len(aggs) == 1
        agg = aggs[0]
        assert agg.item_key == "code:C1"  # 历史映射保持
        assert agg.cum_qty != D("1001")
        assert agg.cum_qty is None, "未证明单位兼容前禁止输出累计数量"
        assert agg.wavg_price is None, "禁止生成跨单位加权平均单价"
        assert agg.cum_amount == D("2000"), "金额（货币口径）如实累计"
        assert agg.status == "incomparable"
        assert any("单位不一致" in w and "不可比" in w for w in agg.warnings)
    finally:
        conn.close()


def test_cross_period_unit_conflict_keeps_consistent_period_values(tmp_path):
    """各期内部单位一致时，期内数量/单价保留（各自口径成立），仅组级封禁。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "C1", "清单A", "", "千克", "1000", "1", "1000"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        pp_values = sorted(
            (pp["qty"], pp["wavg_price"]) for pp in agg.per_period.values()
        )
        assert pp_values == [(D(1), D("1000")), (D(1000), D(1))]
    finally:
        conn.close()


# ---------- 同期混单位 ----------


def test_same_period_mixed_units_blocked(tmp_path):
    """同一期内混单位：该期数量也不得输出相加值。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(1, "downward", "C1", "清单A", "", "千克", "1000", "1", "1000"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.status == "incomparable"
        assert agg.cum_qty is None
        assert agg.wavg_price is None
        for pp in agg.per_period.values():
            assert pp["qty"] is None, "同期混单位数量不可输出"
            assert pp["wavg_price"] is None
        assert any("第1期" in w and "单位不一致" in w for w in agg.warnings)
    finally:
        conn.close()


# ---------- 正例：不得破坏正常资料 ----------


def test_same_unit_positive_case_unchanged(tmp_path):
    """同单位同口径：累计与加权平均单价照常，status 保持 ok。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "C1", "清单A", "", "吨", "2", "1000", "2000"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.status == "ok"
        assert agg.cum_qty == D("3")
        assert agg.cum_amount == D("3000")
        assert agg.wavg_price == D("1000")
    finally:
        conn.close()


def test_alias_units_are_compatible(tmp_path):
    """'吨' 与 'T' 归一化后同单位，不制造假不可比。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "C1", "清单A", "", "T", "1", "1000", "1000"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.status == "ok"
        assert agg.cum_qty == D("2")
        assert agg.wavg_price == D("1000")
    finally:
        conn.close()


# ---------- 空单位的 fail-closed 边界 ----------


def test_empty_unit_with_named_unit_is_incomparable(tmp_path):
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "C1", "清单A", "", "", "1", "1000", "1000"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.status == "incomparable"
        assert agg.cum_qty is None
    finally:
        conn.close()


def test_all_units_empty_keeps_historical_behavior(tmp_path):
    """全部未注明单位：沿用历史口径，不在累计层制造不可比。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "", "1", "1000", "1000"),
        _row(2, "downward", "C1", "清单A", "", "", "1", "1000", "1000"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.status == "ok"
        assert agg.cum_qty == D("2")
        assert agg.wavg_price == D("1000")
    finally:
        conn.close()


# ---------- 无编码同名归组 ----------


def test_same_name_no_code_different_units(tmp_path):
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "", "清单A", "", "千克", "1000", "1", "1000"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.item_key == "name:清单A"
        assert agg.status == "incomparable"
        assert agg.cum_qty is None
    finally:
        conn.close()


# ---------- 同名不同规格 / 不同范围反例 ----------


def test_same_code_different_feature_needs_review(tmp_path):
    """同码同名不同规格：不静默确认合并，降级待复核并警示。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "C30混凝土", "m3", "10", "400", "4000"),
        _row(2, "downward", "C1", "清单A", "C25混凝土", "m3", "10", "380", "3800"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.status == "incomplete"
        assert any("规格不一致" in w and "待人工复核" in w for w in agg.warnings)
        # 单位一致时数量可累计，但组已标记待复核，不作为已确认合并结论
        assert agg.cum_qty == D("20")
    finally:
        conn.close()


def test_same_feature_positive_case_not_flagged(tmp_path):
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "C30混凝土", "m3", "10", "400", "4000"),
        _row(2, "downward", "C1", "清单A", "C30混凝土", "m3", "10", "400", "4000"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.status == "ok"
        assert not any("规格不一致" in w for w in agg.warnings)
    finally:
        conn.close()


def test_period_coverage_gap_still_incomplete(tmp_path):
    """不同范围反例：某期缺该清单 → 漏项警示（既有行为回归）。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "C2", "清单B", "", "吨", "1", "500", "500"),
    ])
    try:
        aggs = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )
        by_key = {a.item_key: a for a in aggs}
        a = by_key["code:C1"]
        assert a.status == "incomplete"
        assert any("第2期无此清单" in w and "漏项" in w for w in a.warnings)
        assert a.cum_qty == D("1"), "缺失期不补 0"
    finally:
        conn.close()


# ---------- 持久化与传播 ----------


def test_persist_period_totals_nulls_conflicted_period_qty(tmp_path):
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(1, "downward", "C1", "清单A", "", "千克", "1000", "1", "1000"),
    ])
    try:
        aggs = aggregate.aggregate_project(conn, info.project_id, direction="downward")
        aggregate.persist_period_totals(conn, info.project_id, aggs)
        rows = conn.execute(
            """SELECT pt.qty_sum, pt.wavg_price, pt.effective_amount_sum
               FROM period_totals pt
               JOIN settlement_periods sp ON sp.id = pt.period_id
               WHERE pt.project_id=? AND sp.period_no=1""",
            (info.project_id,),
        ).fetchall()
        assert len(rows) == 1
        assert rows[0]["qty_sum"] is None, "同期混单位数量禁止入 period_totals"
        assert rows[0]["wavg_price"] is None
        assert rows[0]["effective_amount_sum"] == "2000"
    finally:
        conn.close()


def test_summary_counts_incomparable_group(tmp_path):
    """不可比状态向项目摘要（界面数据源）传播。"""
    from jiadun.core.reporting import build_project_summary

    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "C1", "清单A", "", "千克", "1000", "1", "1000"),
    ])
    try:
        summary = build_project_summary(conn, info.project_id)
        amounts = summary.amounts["downward"]
        assert amounts["incomparable_groups"] == 1
        assert amounts["ok_groups"] == 0
    finally:
        conn.close()


def test_export_settlement_summary_marks_incomparable(tmp_path):
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "downward", "C1", "清单A", "", "千克", "1000", "1", "1000"),
    ])
    try:
        import openpyxl

        from jiadun.core.export import excel_export

        wb = openpyxl.Workbook()
        excel_export.export_settlement_summary(
            conn, info.project_id, wb, direction="downward"
        )
        ws = wb["对下结算累计表"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        data_row = next(r for r in rows if r[0] == "C1")
        # 列：0编码 1名称 2单位 3/4各期金额 5累计数量 6累计金额 7加权单价 8状态
        assert data_row[5] is None, "累计数量列不得输出 1001"
        assert data_row[6] == D("2000")
        assert data_row[8] and "不可比" in str(data_row[8])
    finally:
        conn.close()


def test_export_updown_comparison_discloses_unit_conflict(tmp_path):
    """同方向多期混单位：该侧数量封禁，口径说明列显式披露。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "upward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(2, "upward", "C1", "清单A", "", "千克", "1000", "1", "1000"),
        _row(1, "downward", "C1", "清单A", "", "吨", "2", "1000", "2000"),
    ])
    try:
        import openpyxl

        from jiadun.core.export import excel_export

        wb = openpyxl.Workbook()
        excel_export.export_updown_comparison(conn, info.project_id, wb)
        ws = wb["对上对下对比表"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        data_row = next(r for r in rows if r[0] == "C1")
        # 对上侧跨期混单位 → 数量封禁；对下侧同单位 → 数量照常
        assert data_row[2] is None, "对上侧混单位数量不得输出 1001"
        assert data_row[3] == D("2")
        note = str(data_row[7] or "")
        assert "对上结算" in note and "单位不一致" in note and "不可比" in note
    finally:
        conn.close()


def test_export_updown_comparison_same_unit_positive(tmp_path):
    """同单位正例：对上对下数量照常导出，不出现假不可比。"""
    info, conn = _build_project(tmp_path, [
        _row(1, "upward", "C1", "清单A", "", "吨", "1", "1000", "1000"),
        _row(1, "downward", "C1", "清单A", "", "吨", "2", "1000", "2000"),
    ])
    try:
        import openpyxl

        from jiadun.core.export import excel_export

        wb = openpyxl.Workbook()
        excel_export.export_updown_comparison(conn, info.project_id, wb)
        ws = wb["对上对下对比表"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        data_row = next(r for r in rows if r[0] == "C1")
        assert data_row[2] == D("1")
        assert data_row[3] == D("2")
        assert "单位不一致" not in str(data_row[7] or "")
    finally:
        conn.close()


@pytest.mark.parametrize("unit_a,unit_b", [("吨", "千克"), ("m3", "m2"), ("个", "樘")])
def test_various_incompatible_unit_pairs(tmp_path, unit_a, unit_b):
    info, conn = _build_project(tmp_path, [
        _row(1, "downward", "C1", "清单A", "", unit_a, "1", "100", "100"),
        _row(2, "downward", "C1", "清单A", "", unit_b, "1", "100", "100"),
    ])
    try:
        agg = aggregate.aggregate_project(
            conn, info.project_id, direction="downward"
        )[0]
        assert agg.status == "incomparable"
        assert agg.cum_qty is None
    finally:
        conn.close()
