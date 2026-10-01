"""标准描述、特征和计量口径一致的清单，不能仅因来源编码不同而漏配。"""
from pathlib import Path

import pytest

from jiadun.core.engine import quantity_control as qc
from jiadun.core.models import project as pm


@pytest.fixture(autouse=True)
def _settings_sandbox(tmp_path, monkeypatch):
    """create_project 的工作区登记不触真实 settings（临时 settings.json）。"""
    sandbox = tmp_path / "settings-sandbox"
    sandbox.mkdir(exist_ok=True)
    settings_file = sandbox / "settings.json"
    settings_file.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(pm, "_SETTINGS_FILE", settings_file)
    yield


def test_exact_description_and_feature_merge_across_source_codes(tmp_path):
    info = pm.create_project('跨来源编码描述匹配', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        for key, direction, kind, code, quantity in [
            ('U', 'upward', 'contract', 'C1', '100'),
            ('A', 'downward', 'settlement', 'D9', '60'),
            ('B', 'downward', 'settlement', None, '60'),
        ]:
            context = qc.create_period_context(
                conn, info.project_id, direction=direction, contract_key=key,
                business_period_no=1, unit_name=key, doc_kind=kind,
                work_scope='C30混凝土材料供应', actor='user', reason='同实体同计量工作范围',
            )
            conn.execute(
                'INSERT INTO line_items(period_id,code,name,feature,unit,quantity) VALUES (?,?,?,?,?,?)',
                (context['period_id'], code, '现浇混凝土', 'C30', 'm3', quantity),
            )
            qc.confirm_period_context(conn, info.project_id, context['context_id'],
                                      actor='user', reason='确认同对象的完整标准描述和工作范围')
        controls = [c for c in qc.build_quantity_ledger(conn, info.project_id)['quantity_controls']
                    if c['building'] is None]
        assert len(controls) == 1
        assert controls[0]['downstream_quantity'] == '120'
        assert controls[0]['delta_vs_contract'] == '20'
        assert controls[0]['status'] == 'FAIL'
    finally:
        conn.close()


def _context_with_line(conn, pid, *, direction, key, kind, code, name, feature,
                       unit, quantity, scope):
    context = qc.create_period_context(
        conn, pid, direction=direction, contract_key=key, business_period_no=1,
        unit_name=key, doc_kind=kind, work_scope=scope, actor='user',
        reason='登记')
    conn.execute(
        'INSERT INTO line_items(period_id,code,name,feature,unit,quantity)'
        ' VALUES (?,?,?,?,?,?)',
        (context['period_id'], code, name, feature, unit, quantity))
    qc.confirm_period_context(conn, pid, context['context_id'], actor='user',
                              reason='确认')
    return context


def _controls_by_desc(conn, pid, *, scope='材料供应'):
    return [c for c in qc.build_quantity_ledger(conn, pid)['quantity_controls']
            if c['work_scope'] == scope and c['building'] is None]


def test_same_code_different_names_not_merged(tmp_path):
    """同码异名：描述身份不同，不得自动归并。"""
    info = pm.create_project('同码异名', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        _context_with_line(conn, info.project_id, direction='upward', key='U',
                           kind='contract', code='C1', name='甲项',
                           feature='C30', unit='m3', quantity='100',
                           scope='材料供应')
        _context_with_line(conn, info.project_id, direction='downward', key='A',
                           kind='settlement', code='C1', name='乙项',
                           feature='C30', unit='m3', quantity='60',
                           scope='材料供应')
        assert len(_controls_by_desc(conn, info.project_id)) == 2
    finally:
        conn.close()


def test_different_features_not_merged(tmp_path):
    """C30/C35 跨特征：不归并。"""
    info = pm.create_project('跨特征', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        _context_with_line(conn, info.project_id, direction='upward', key='U',
                           kind='contract', code='C1', name='现浇混凝土',
                           feature='C30', unit='m3', quantity='100',
                           scope='材料供应')
        _context_with_line(conn, info.project_id, direction='downward', key='A',
                           kind='settlement', code='C1', name='现浇混凝土',
                           feature='C35', unit='m3', quantity='60',
                           scope='材料供应')
        assert len(_controls_by_desc(conn, info.project_id)) == 2
    finally:
        conn.close()


def test_empty_feature_keeps_conservative_code_identity(tmp_path):
    """空特征：不同编码保守待确认，不按名称归并。"""
    info = pm.create_project('空特征', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        _context_with_line(conn, info.project_id, direction='upward', key='U',
                           kind='contract', code='C1', name='现浇混凝土',
                           feature='', unit='m3', quantity='100',
                           scope='材料供应')
        _context_with_line(conn, info.project_id, direction='downward', key='A',
                           kind='settlement', code='D9', name='现浇混凝土',
                           feature='', unit='m3', quantity='60',
                           scope='材料供应')
        assert len(_controls_by_desc(conn, info.project_id)) == 2
    finally:
        conn.close()


def test_different_scope_or_unit_not_merged(tmp_path):
    """不同口径或不同量纲：不归并。"""
    info = pm.create_project('口径量纲', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        _context_with_line(conn, info.project_id, direction='upward', key='U',
                           kind='contract', code='C1', name='现浇混凝土',
                           feature='C30', unit='m3', quantity='100',
                           scope='材料供应')
        _context_with_line(conn, info.project_id, direction='downward', key='A',
                           kind='settlement', code='C1', name='现浇混凝土',
                           feature='C30', unit='m3', quantity='60',
                           scope='浇筑劳务')
        _context_with_line(conn, info.project_id, direction='downward', key='B',
                           kind='settlement', code='C1', name='现浇混凝土',
                           feature='C30', unit='项', quantity='60',
                           scope='材料供应')
        controls = qc.build_quantity_ledger(conn, info.project_id)['quantity_controls']
        material = [c for c in controls if c['work_scope'] == '材料供应'
                    and c['building'] is None]
        labor = [c for c in controls if c['work_scope'] == '浇筑劳务'
                 and c['building'] is None]
        assert len(material) == 2  # m3 与"项"量纲不同
        assert len(labor) == 1
    finally:
        conn.close()


def test_manual_standard_key_not_auto_merged(tmp_path):
    """人工标准键优先：不同键不因描述相同而自动归并。"""
    info = pm.create_project('人工键', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        for key, std in (('U', 'STD-A'), ('A', 'STD-B')):
            context = _context_with_line(
                conn, info.project_id, direction='upward' if key == 'U'
                else 'downward', key=key,
                kind='contract' if key == 'U' else 'settlement', code='C1',
                name='现浇混凝土', feature='C30', unit='m3',
                quantity='100' if key == 'U' else '60', scope='材料供应')
            line_id = conn.execute(
                'SELECT id FROM line_items WHERE period_id=?',
                (context['period_id'],)).fetchone()['id']
            qc.set_line_context(conn, info.project_id, line_id,
                                standard_key=std, actor='user', reason='人工键')
            qc.confirm_line_context(conn, info.project_id, line_id,
                                    actor='user', reason='确认')
        controls = [c for c in qc.build_quantity_ledger(
            conn, info.project_id)['quantity_controls']
            if c['building'] is None and c['work_scope'] == '材料供应']
        assert len(controls) == 2
        assert all(c['standard_key'] in ('STD-A', 'STD-B') for c in controls)
    finally:
        conn.close()


def test_merge_keeps_source_codes_and_rule_status(tmp_path):
    """归并组保留各原编码来源；规则状态 auto:desc 不冒记人工。"""
    info = pm.create_project('来源保留', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        _context_with_line(conn, info.project_id, direction='upward', key='U',
                           kind='contract', code='C1', name='现浇混凝土',
                           feature='C30', unit='m3', quantity='100',
                           scope='材料供应')
        _context_with_line(conn, info.project_id, direction='downward', key='A',
                           kind='settlement', code='D9', name='现浇混凝土',
                           feature='C30', unit='m3', quantity='60',
                           scope='材料供应')
        control = _controls_by_desc(conn, info.project_id)[0]
        assert control['source_codes'] == ['C1', 'D9']
        assert control['identity_source'] == 'auto:desc'
        assert control['status'] == 'PASS'
        assert control['delta_vs_contract'] == '-40'
    finally:
        conn.close()


def test_pending_coverage_still_blocks_merged_group(tmp_path):
    """归并组存在待确认行（异码/无码同名）时仍被覆盖拦截。"""
    info = pm.create_project('归并覆盖', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        _context_with_line(conn, info.project_id, direction='upward', key='U',
                           kind='contract', code='C1', name='现浇混凝土',
                           feature='C30', unit='m3', quantity='100',
                           scope='材料供应')
        _context_with_line(conn, info.project_id, direction='downward', key='A',
                           kind='settlement', code='D9', name='现浇混凝土',
                           feature='C30', unit='m3', quantity='60',
                           scope='材料供应')
        pending = qc.create_period_context(
            conn, info.project_id, direction='downward', contract_key='B',
            business_period_no=1, unit_name='B', doc_kind='settlement',
            work_scope='材料供应', actor='user', reason='待确认期次')
        conn.execute(
            'INSERT INTO line_items(period_id,code,name,feature,unit,quantity)'
            ' VALUES (?,?,?,?,?,?)',
            (pending['period_id'], None, '现浇混凝土', 'C30', 'm3', '60'))
        control = _controls_by_desc(conn, info.project_id)[0]
        assert control['status'] == 'PENDING'
        assert control['status_vs_contract'] == 'PENDING'
        assert control['excluded_details']
    finally:
        conn.close()


def test_same_contract_comparison_merges_across_source_codes(tmp_path):
    """同合同比较也走描述身份：U 合同 C1=100 vs 同 U 结算 D9=60。

    名称/特征/量纲/口径一致时，比较层不得拆成两个 PENDING，应得出
    "按合同未超"（结余 40）。
    """
    info = pm.create_project('同合同比较归并', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        for kind, code, quantity in (('contract', 'C1', '100'),
                                     ('settlement', 'D9', '60')):
            context = qc.create_period_context(
                conn, info.project_id, direction='upward', contract_key='U',
                business_period_no=1, unit_name='U', doc_kind=kind,
                work_scope='材料供应', actor='user', reason='登记')
            conn.execute(
                'INSERT INTO line_items(period_id,code,name,feature,unit,quantity)'
                ' VALUES (?,?,?,?,?,?)',
                (context['period_id'], code, '现浇混凝土', 'C30', 'm3', quantity))
            qc.confirm_period_context(conn, info.project_id,
                                      context['context_id'], actor='user',
                                      reason='确认')
        ledger = qc.build_quantity_ledger(conn, info.project_id)
        comparisons = [c for c in ledger['comparisons']
                       if c['building'] is None and c['work_scope'] == '材料供应']
        assert len(comparisons) == 1, comparisons
        assert comparisons[0]['contract_quantity'] == '100'
        assert comparisons[0]['settlement_quantity'] == '60'
        assert comparisons[0]['delta'] == '-40'
        assert comparisons[0]['status'] == 'PASS'
        # 来源保留：比较两侧原编码可见
        assert {s['original_code'] for item in ledger['items']
                for s in item['sources']} == {'C1', 'D9'}
    finally:
        conn.close()


def test_same_contract_multi_code_rows_one_upstream_baseline(tmp_path):
    """同一合同一期内 C1=40 + C2=60 同描述 → 单一上游基准 100。

    A 结算 D9=120：上限 100、超 20 FAIL，不得因两条 upward items 触发
    "多个上游基准"冲突。
    """
    info = pm.create_project('单基准多编码', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        context = qc.create_period_context(
            conn, info.project_id, direction='upward', contract_key='U',
            business_period_no=1, unit_name='U', doc_kind='contract',
            work_scope='材料供应', actor='user', reason='登记')
        for code, quantity in (('C1', '40'), ('C2', '60')):
            conn.execute(
                'INSERT INTO line_items(period_id,code,name,feature,unit,quantity)'
                ' VALUES (?,?,?,?,?,?)',
                (context['period_id'], code, '现浇混凝土', 'C30', 'm3', quantity))
        qc.confirm_period_context(conn, info.project_id, context['context_id'],
                                  actor='user', reason='确认')
        _context_with_line(conn, info.project_id, direction='downward', key='A',
                           kind='settlement', code='D9', name='现浇混凝土',
                           feature='C30', unit='m3', quantity='120',
                           scope='材料供应')
        ledger = qc.build_quantity_ledger(conn, info.project_id)
        controls = [c for c in ledger['quantity_controls']
                    if c['building'] is None and c['work_scope'] == '材料供应']
        assert len(controls) == 1
        control = controls[0]
        assert control['baselines']['upward_contract']['quantity'] == '100'
        assert control['downstream_quantity'] == '120'
        assert control['delta_vs_contract'] == '20'
        assert control['status_vs_contract'] == 'FAIL'
        assert control['status'] == 'FAIL'
        assert control['baselines']['upward_contract']['contributors'] == 1
        assert '不得默认叠加' not in control['reason']
        # 组内原编码来源保留（C1/C2 为 U 合同行，D9 为 A 结算行）
        assert set(control['source_codes']) == {'C1', 'C2', 'D9'}
        item = next(x for x in ledger['items']
                    if x['work_scope'] == '材料供应'
                    and x['direction'] == 'upward')
        assert item['identity_status'] == 'rule_accepted_desc'
        assert '描述规则自动接受' in item['identity_basis']
        assert sorted(item['source_codes']) == ['C1', 'C2']
    finally:
        conn.close()


def test_unrelated_incomplete_pending_row_does_not_block_description():
    identity = qc._line_identity("C1", "混凝土", "C30", None, "m3")
    entry = {"identity": qc._line_identity("X9", "防水", "", None, "m3"),
             "code": "X9", "name_key": "防水", "feature": "", "base": "m3"}
    assert not qc._identity_may_match(identity, "C1", "混凝土", entry)


def test_source_export_preserves_codes_and_rule_and_human_states(tmp_path):
    from openpyxl import Workbook, load_workbook

    from jiadun.core.export.excel_export import export_quantity_control_sheets

    info = pm.create_project('导出规则与人工来源', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        context = _context_with_line(
            conn, info.project_id, direction='upward', key='U', kind='contract',
            code='C1', name='混凝土', feature='C30', unit='m3', quantity='100', scope='材料')
        conn.execute('UPDATE quantity_period_context SET building_scope=?,building_status=? WHERE id=?',
                     ('1号楼', 'confirmed', context['context_id']))
        ledger = qc.build_quantity_ledger(conn, info.project_id)
        workbook = Workbook()
        export_quantity_control_sheets(ledger, workbook)
        path = tmp_path / 'quantity.xlsx'
        workbook.save(path)
        reopened = load_workbook(path)
        data = list(reopened['有效数量与来源'].values)
        row = dict(zip(data[0], data[1], strict=True))
        assert row['原编码'] == 'C1'
        assert row['原名称'] == '混凝土'
        assert row['原项目特征'] == 'C30'
        assert row['归并状态'] == '规则自动接受'
        assert '描述规则自动接受' in row['归并依据']
        assert row['楼栋识别状态'] == '人工期次确认'
        assert row['楼栋识别依据'] == '期次确认范围'
        assert row['标准行数量'] == '100'
        reopened.close()
    finally:
        conn.close()


def test_pending_original_code_blocks_merged_object_with_multiple_codes(tmp_path):
    info = pm.create_project('归并原编码覆盖', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        for key, direction, kind, code, quantity in [
            ('U', 'upward', 'contract', 'C1', '100'),
            ('A', 'downward', 'settlement', 'D9', '60'),
        ]:
            _context_with_line(conn, info.project_id, direction=direction, key=key,
                               kind=kind, code=code, name='混凝土', feature='C30',
                               unit='m3', quantity=quantity, scope='材料供应')
        pending = qc.create_period_context(
            conn, info.project_id, direction='downward', contract_key='B',
            business_period_no=1, unit_name='B', doc_kind='settlement',
            work_scope='材料供应', actor='user', reason='待补描述')
        line_id = conn.execute(
            'INSERT INTO line_items(period_id,code,name,feature,unit,quantity) VALUES (?,?,?,?,?,?)',
            (pending['period_id'], 'D9', '', '', 'm3', None)).lastrowid
        control = _controls_by_desc(conn, info.project_id)[0]
        assert control['status_vs_contract'] == 'PENDING'
        assert any(row['line_item_id'] == line_id for row in control['excluded_details'])
    finally:
        conn.close()
