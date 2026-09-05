"""结论性框架测试：对上控制基准结论持久化 + 费率基数解析/试算 + 报告导出。

宪章约束：五态如实呈现不强行 PASS；税口径词表归一（included/incl_tax 同义）；
缺失金额行不得按 0 参与合计；多候选并存不得自动挑选；结论/试算只追加不覆盖。
"""
from __future__ import annotations

import docx as docx_lib
import pytest
from openpyxl import Workbook
from tests.unit.test_contract_extract import SAMPLE_CONTRACT

from jiadun.core import document_intake
from jiadun.core.contracts import extract, rate_rules
from jiadun.core.db import migrations
from jiadun.core.engine import control_baseline as cb
from jiadun.core.export import conclusions_report, excel_export
from jiadun.core.models.source_file import import_file


@pytest.fixture()
def project_db(tmp_path):
    db_path = tmp_path / "project.db"
    migrations.migrate(db_path, tmp_path / "backups")
    conn = migrations.connect(db_path)
    with conn:
        project_id = conn.execute(
            """INSERT INTO projects(name, schema_version, workspace_path, created_at)
               VALUES (?,?,?,?)""",
            ("结论框架测试", migrations.LATEST_SCHEMA_VERSION, str(tmp_path), "2026"),
        ).lastrowid
    yield conn, int(project_id), tmp_path
    conn.close()


def make_period(conn, pid, period_no, title, direction, tax_mode="unknown"):
    with conn:
        cur = conn.execute(
            """INSERT INTO settlement_periods(
                   project_id, period_no, title, direction, tax_mode)
               VALUES (?,?,?,?,?)""",
            (pid, period_no, title, direction, tax_mode),
        )
        return int(cur.lastrowid)


def add_item(conn, period_id, name, amount):
    with conn:
        conn.execute(
            "INSERT INTO line_items(period_id, name, amount) VALUES (?,?,?)",
            (period_id, name, amount),
        )


def confirmed_baseline(conn, pid, amount, *, tax_basis="incl_tax"):
    baseline_id = cb.create_candidate_manual(
        conn, pid, amount, source_note="终审报告审定金额", tax_basis=tax_basis)
    cb.set_baseline_review(
        conn, pid, baseline_id, "confirmed", reason="终审报告核对一致")
    return baseline_id


def confirmed_framework_rule(conn, pid, pdir, *, rate_text=None, base_type, tax_basis,
                             cap=None, floor=None):
    """写入含费率条款的框架协议 → 扫描候选 → 人工确认为结构化规则。"""
    content = rate_text or "协作方按对上结算价的3%上缴管理费。"
    src = pdir / f"框架协议_{base_type}_{tax_basis}.txt"
    src.write_text(content, encoding="utf-8")
    sf = import_file(conn, pid, pdir, src)
    document_intake.record_document(
        conn, pid, sf.file_id, category="upward_framework_management",
        parse_status="evidence_only", detail="", parser="")
    rate_rules.import_rate_candidates(conn, pid, sf.file_id)
    rules = rate_rules.list_rate_rules(conn, pid, status="candidate")
    assert rules, "扫描必须产出候选"
    rule_id = int(rules[0]["id"])
    rate_rules.confirm_rate_rule(
        conn, pid, rule_id,
        base_type=base_type,
        base_definition="测试用结构化基数",
        tax_basis=tax_basis,
        cap=cap, floor=floor,
    )
    return rule_id


class TestSchemaV54:
    def test_new_tables_exist(self, project_db):
        conn, pid, _ = project_db
        names = {
            r["name"] for r in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            ).fetchall()
        }
        assert {"control_conclusions", "rate_rule_applications"} <= names
        row = conn.execute(
            "SELECT schema_version FROM projects WHERE id=?", (pid,)).fetchone()
        assert row["schema_version"] == 54


class TestTaxVocabularyNormalization:
    def test_same_meaning_vocab_not_incomparable(self, project_db):
        """回归（红→绿）：基准 included（旧词表）vs 期次 incl_tax 不得误判 INCOMPARABLE。"""
        conn, pid, _ = project_db
        # 旧数据形态：基准直接写 included；期次 tax_mode 是 incl_tax
        with conn:
            conn.execute(
                "UPDATE control_baselines SET tax_basis='included' WHERE id=?",
                (confirmed_baseline(conn, pid, "1200000"),))
        period_id = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        baseline_id = int(conn.execute(
            "SELECT id FROM control_baselines WHERE project_id=?", (pid,)
        ).fetchone()["id"])
        result = cb.compare_upward_result(
            conn, pid, baseline_id, "1100000",
            settlement_tax_basis="incl_tax", period_id=period_id)
        assert result["status"] == "PASS", (
            "同义税口径（included/incl_tax）被误判 INCOMPARABLE 是词表缺陷"
        )

    def test_real_mismatch_still_incomparable(self, project_db):
        conn, pid, _ = project_db
        baseline_id = confirmed_baseline(conn, pid, "1200000", tax_basis="incl_tax")
        period_id = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="excl_tax")
        result = cb.compare_upward_result(
            conn, pid, baseline_id, "1100000",
            settlement_tax_basis="excl_tax", period_id=period_id)
        assert result["status"] == "INCOMPARABLE"

    def test_candidate_registration_normalizes_legacy_vocab(self, project_db):
        conn, pid, _ = project_db
        baseline_id = cb.create_candidate_manual(
            conn, pid, "100", source_note="测试", tax_basis="included")
        stored = conn.execute(
            "SELECT tax_basis FROM control_baselines WHERE id=?", (baseline_id,)
        ).fetchone()["tax_basis"]
        assert stored == "incl_tax"


class TestPeriodTaxModeReview:
    def test_confirm_writes_evidence(self, project_db):
        conn, pid, _ = project_db
        period_id = make_period(conn, pid, 1, "对上第1期", "upward")
        result = cb.set_period_tax_mode(
            conn, pid, period_id, "excl_tax", reason="表头注明不含税单价")
        assert result["before"] == "unknown"
        assert result["tax_mode"] == "excl_tax"
        kinds = {
            r["kind"] for r in conn.execute(
                "SELECT kind FROM evidence WHERE project_id=?", (pid,)).fetchall()
        }
        assert "period_tax_mode_review" in kinds

    def test_invalid_value_rejected(self, project_db):
        conn, pid, _ = project_db
        period_id = make_period(conn, pid, 1, "对上第1期", "upward")
        with pytest.raises(ValueError, match="税口径"):
            cb.set_period_tax_mode(conn, pid, period_id, "maybe", reason="x")
        stored = conn.execute(
            "SELECT tax_mode FROM settlement_periods WHERE id=?", (period_id,)
        ).fetchone()["tax_mode"]
        assert stored == "unknown"

    def test_reason_required(self, project_db):
        conn, pid, _ = project_db
        period_id = make_period(conn, pid, 1, "对上第1期", "upward")
        with pytest.raises(ValueError, match="依据"):
            cb.set_period_tax_mode(conn, pid, period_id, "incl_tax", reason="  ")

    def test_wrong_project_rejected(self, project_db):
        conn, pid, _ = project_db
        with pytest.raises(ValueError, match="不属于当前项目"):
            cb.set_period_tax_mode(conn, pid, 99999, "incl_tax", reason="x")


class TestConclusionPersistence:
    def test_compare_persists_conclusion_snapshot(self, project_db):
        conn, pid, _ = project_db
        baseline_id = confirmed_baseline(conn, pid, "1200000")
        period_id = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        result = cb.compare_upward_result(
            conn, pid, baseline_id, "1300000",
            settlement_tax_basis="incl_tax", period_id=period_id)
        assert result["status"] == "FAIL"
        rows = cb.list_control_conclusions(conn, pid)
        assert len(rows) == 1
        assert rows[0]["id"] == result["conclusion_id"]
        assert rows[0]["period_id"] == period_id
        assert rows[0]["status"] == "FAIL"
        assert rows[0]["delta"] == "100000"
        assert rows[0]["evidence_id"] == result["evidence_id"]
        assert rows[0]["period_no"] == 1

    def test_history_append_only(self, project_db):
        conn, pid, _ = project_db
        baseline_id = confirmed_baseline(conn, pid, "1200000")
        make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        cb.compare_upward_result(
            conn, pid, baseline_id, "1100000", settlement_tax_basis="incl_tax")
        cb.compare_upward_result(
            conn, pid, baseline_id, "1300000", settlement_tax_basis="incl_tax")
        rows = cb.list_control_conclusions(conn, pid)
        assert len(rows) == 2
        assert [r["status"] for r in rows] == ["FAIL", "PASS"]  # 倒序
        assert rows[0]["id"] > rows[1]["id"], "旧结论必须保留（不覆盖）"

    def test_period_binding_validated(self, project_db):
        conn, pid, _ = project_db
        baseline_id = confirmed_baseline(conn, pid, "1200000")
        with pytest.raises(ValueError, match="对上结算期次"):
            cb.compare_upward_result(
                conn, pid, baseline_id, "1",
                settlement_tax_basis="incl_tax", period_id=99999)


class TestRateBaseResolution:
    def test_upward_periods_sum_decimal_exact(self, project_db):
        conn, pid, pdir = project_db
        p1 = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        p2 = make_period(conn, pid, 2, "对上第2期", "upward", tax_mode="incl_tax")
        add_item(conn, p1, "明细A", "1000.50")
        add_item(conn, p2, "明细B", "2000.25")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount", tax_basis="incl_tax")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_RESOLVED
        assert result["base_amount"] == "3000.75"
        assert [p["period_no"] for p in result["breakdown"]["periods"]] == [1, 2]

    def test_period_scope_limits_base(self, project_db):
        conn, pid, pdir = project_db
        p1 = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        p2 = make_period(conn, pid, 2, "对上第2期", "upward", tax_mode="incl_tax")
        add_item(conn, p1, "明细A", "1000.50")
        add_item(conn, p2, "明细B", "2000.25")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount", tax_basis="incl_tax")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id, period_id=p1)
        assert result["status"] == rate_rules.BASE_RESOLVED
        assert result["base_amount"] == "1000.50"

    def test_missing_amount_rows_block_total(self, project_db):
        """缺失金额行不得按 0 参与合计（宪章：缺失≠0）。"""
        conn, pid, pdir = project_db
        p1 = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        add_item(conn, p1, "明细A", "1000.50")
        add_item(conn, p1, "缺金额行", None)
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount", tax_basis="incl_tax")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_PENDING
        assert "金额缺失" in result["reason"]

    def test_unknown_period_tax_pending(self, project_db):
        conn, pid, pdir = project_db
        p1 = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="unknown")
        add_item(conn, p1, "明细A", "100")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount", tax_basis="incl_tax")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_PENDING
        assert "税口径未确认" in result["reason"]

    def test_mixed_period_tax_incomparable(self, project_db):
        conn, pid, pdir = project_db
        make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        make_period(conn, pid, 2, "对上第2期", "upward", tax_mode="excl_tax")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount", tax_basis="incl_tax")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_INCOMPARABLE
        assert "混用" in result["reason"]

    def test_rule_period_tax_mismatch_incomparable(self, project_db):
        conn, pid, pdir = project_db
        make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        add_item(conn, conn.execute(
            "SELECT id FROM settlement_periods WHERE project_id=? AND period_no=1",
            (pid,)).fetchone()["id"], "明细A", "100")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount", tax_basis="excl_tax")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_INCOMPARABLE

    def test_excl_tax_base_never_converted_from_incl(self, project_db):
        """不含税基数不得从含税合计按猜测税率换算。"""
        conn, pid, pdir = project_db
        make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        add_item(conn, conn.execute(
            "SELECT id FROM settlement_periods WHERE project_id=? AND period_no=1",
            (pid,)).fetchone()["id"], "明细A", "11300")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_excl_tax",
            tax_basis="excl_tax")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_INCOMPARABLE
        assert "换算" in result["reason"]

    def test_no_periods_pending(self, project_db):
        conn, pid, pdir = project_db
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="downward_settlement_amount", tax_basis="incl_tax")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_PENDING
        assert "对下" in result["reason"]

    def test_custom_requires_manual_amount(self, project_db):
        conn, pid, pdir = project_db
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="custom", tax_basis="unknown")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_MANUAL_REQUIRED
        result = rate_rules.resolve_rate_base(conn, pid, rule_id, custom_amount="1234.56")
        assert result["status"] == rate_rules.BASE_RESOLVED
        assert result["base_amount"] == "1234.56"

    def test_contract_amount_single_confirmed_fact_resolved(self, project_db):
        conn, pid, pdir = project_db
        src = pdir / "分包合同.docx"
        d = docx_lib.Document()
        for line in SAMPLE_CONTRACT.splitlines():
            d.add_paragraph(line)
        d.save(str(src))
        extract.import_contract(conn, pid, pdir, src, doc_type="subcontract")
        fact_ids = [
            int(r["id"]) for r in conn.execute(
                """SELECT cf.id FROM contract_facts cf
                   JOIN contract_docs cd ON cd.id=cf.doc_id
                   WHERE cd.project_id=? AND cf.fact_key='contract_amount'
                     AND cf.fact_value LIKE '%1286500%'""",
                (pid,)).fetchall()
        ]
        assert fact_ids, "样例合同必须抽出签约合同价金额事实"
        extract.set_fact_review(conn, pid, fact_ids[0], "confirmed", reason="核对一致")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="contract_amount", tax_basis="unknown")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_RESOLVED
        assert result["base_amount"] == "1286500.00"
        assert result["breakdown"]["fact_id"] == fact_ids[0]

    def test_confirmed_none_valued_fact_excluded_not_crash(self, project_db):
        """空值占位事实被确认后：不得参与基数，也不得让解析崩溃。"""
        conn, pid, pdir = project_db
        src = pdir / "空值合同.docx"
        d = docx_lib.Document()
        d.add_paragraph("第二条 合同价款")
        d.add_paragraph("（本页未填写金额）")
        d.save(str(src))
        extract.import_contract(conn, pid, pdir, src, doc_type="subcontract")
        none_facts = [
            int(r["id"]) for r in conn.execute(
                """SELECT cf.id FROM contract_facts cf
                   JOIN contract_docs cd ON cd.id=cf.doc_id
                   WHERE cd.project_id=? AND cf.fact_key='contract_amount'
                     AND cf.fact_value IS NULL""",
                (pid,)).fetchall()
        ]
        assert none_facts, "空值合同应抽出 None 占位事实"
        extract.set_fact_review(conn, pid, none_facts[0], "confirmed", reason="误操作确认")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="contract_amount", tax_basis="unknown")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_PENDING
        assert any(f["reason"] == "非可解析金额"
                   for f in result["breakdown"]["excluded_non_amount_facts"])

    def test_contract_amount_zero_facts_pending(self, project_db):
        conn, pid, pdir = project_db
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="contract_amount", tax_basis="unknown")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_PENDING

    def test_contract_amount_two_confirmed_facts_conflict(self, project_db):
        conn, pid, pdir = project_db
        for name, amount in (("合同A.docx", "1000000"), ("合同B.docx", "2000000")):
            src = pdir / name
            d = docx_lib.Document()
            d.add_paragraph(f"合同价款为人民币{amount}元。")
            d.add_paragraph("费用由管理费按3%计取。")
            d.save(str(src))
            extract.import_contract(conn, pid, pdir, src, doc_type="subcontract")
        fact_ids = [
            int(r["id"]) for r in conn.execute(
                """SELECT cf.id FROM contract_facts cf
                   JOIN contract_docs cd ON cd.id=cf.doc_id
                   WHERE cd.project_id=? AND cf.fact_key='contract_amount'""",
                (pid,)).fetchall()
        ]
        assert len(fact_ids) == 2
        for fid in fact_ids:
            extract.set_fact_review(conn, pid, fid, "confirmed", reason="核对一致")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="contract_amount", tax_basis="unknown")
        result = rate_rules.resolve_rate_base(conn, pid, rule_id)
        assert result["status"] == rate_rules.BASE_CONFLICT
        assert "不自动挑选" in result["reason"]

    def test_candidate_rule_cannot_resolve(self, project_db):
        conn, pid, pdir = project_db
        src = pdir / "框架协议_候选.txt"
        src.write_text("协作方按对上结算价的3%上缴管理费。", encoding="utf-8")
        sf = import_file(conn, pid, pdir, src)
        document_intake.record_document(
            conn, pid, sf.file_id, category="upward_framework_management",
            parse_status="evidence_only", detail="", parser="")
        rate_rules.import_rate_candidates(conn, pid, sf.file_id)
        rule_id = int(rate_rules.list_rate_rules(conn, pid)[0]["id"])
        with pytest.raises(ValueError, match="已确认"):
            rate_rules.resolve_rate_base(conn, pid, rule_id)


class TestApplyRateRuleToSettlement:
    def test_computed_fee_and_persistence(self, project_db):
        conn, pid, pdir = project_db
        p1 = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        add_item(conn, p1, "明细A", "1000.50")
        add_item(conn, p1, "明细B", "2000.25")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount",
            tax_basis="incl_tax", cap="500", floor=None)
        result = rate_rules.apply_rate_rule_to_settlement(conn, pid, rule_id)
        # 3000.75 × 3% = 90.02（round2）< cap 500 → 90.02
        assert result["base_status"] == rate_rules.BASE_RESOLVED
        assert result["base_amount"] == "3000.75"
        assert result["fee"] == "90.02"
        rows = rate_rules.list_rate_applications(conn, pid)
        assert len(rows) == 1
        assert rows[0]["fee"] == "90.02"
        assert rows[0]["evidence_id"] == result["evidence_id"]
        kinds = {
            r["kind"] for r in conn.execute(
                "SELECT kind FROM evidence WHERE project_id=?", (pid,)).fetchall()
        }
        assert "rate_rule_apply_settlement" in kinds

    def test_cap_applied(self, project_db):
        conn, pid, pdir = project_db
        p1 = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        add_item(conn, p1, "明细A", "10000")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount",
            tax_basis="incl_tax", cap="200")
        result = rate_rules.apply_rate_rule_to_settlement(conn, pid, rule_id)
        assert result["fee"] == "200"
        assert result["compute_detail"]["cap_applied"] == "200"

    def test_blocked_attempt_persisted_without_fee(self, project_db):
        """被阻断的试算也要落档（不计算费用），失败尝试是审计事实。"""
        conn, pid, pdir = project_db
        make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="unknown")
        add_item(conn, conn.execute(
            "SELECT id FROM settlement_periods WHERE project_id=? AND period_no=1",
            (pid,)).fetchone()["id"], "明细A", "1000")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount", tax_basis="incl_tax")
        result = rate_rules.apply_rate_rule_to_settlement(conn, pid, rule_id)
        assert result["fee"] is None
        assert result["base_status"] == rate_rules.BASE_PENDING
        rows = rate_rules.list_rate_applications(conn, pid)
        assert len(rows) == 1
        assert rows[0]["fee"] is None
        assert rows[0]["status"] == "pending"

    def test_manual_amount_path_uses_custom(self, project_db):
        conn, pid, pdir = project_db
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="custom", tax_basis="unknown")
        result = rate_rules.apply_rate_rule_to_settlement(
            conn, pid, rule_id, custom_amount="1000")
        assert result["fee"] == "30.00"


class TestReportExport:
    def test_markdown_contains_all_states_and_evidence(self, project_db):
        conn, pid, _ = project_db
        baseline_id = confirmed_baseline(conn, pid, "1200000", tax_basis="incl_tax")
        period_id = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        fail = cb.compare_upward_result(
            conn, pid, baseline_id, "1300000",
            settlement_tax_basis="incl_tax", period_id=period_id)
        inc = cb.compare_upward_result(
            conn, pid, baseline_id, "100",
            settlement_tax_basis="excl_tax", period_id=period_id)
        assert fail["status"] == "FAIL" and inc["status"] == "INCOMPARABLE"
        md = conclusions_report.build_conclusions_markdown(conn, pid)
        assert "对上控制基准结论" in md
        assert "INCOMPARABLE（不可比，不得强行比较）" in md
        assert "FAIL（超基准，不构成违规或责任认定）" in md
        assert str(fail["evidence_id"]) in md and str(inc["evidence_id"]) in md
        assert "不构成违规" in md

    def test_markdown_empty_project_fail_closed_text(self, project_db):
        conn, pid, _ = project_db
        md = conclusions_report.build_conclusions_markdown(conn, pid)
        assert "无结论不等于通过" in md
        assert "无规则不等于免计取" in md

    def test_markdown_includes_rate_applications(self, project_db):
        conn, pid, pdir = project_db
        p1 = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        add_item(conn, p1, "明细A", "1000")
        rule_id = confirmed_framework_rule(
            conn, pid, pdir, base_type="upward_settlement_amount", tax_basis="incl_tax")
        rate_rules.apply_rate_rule_to_settlement(conn, pid, rule_id)
        md = conclusions_report.build_conclusions_markdown(conn, pid)
        assert "框架/管理性协议费率规则与试算" in md
        assert "30.00 元" in md

    def test_excel_sheets_present_with_conclusions(self, project_db):
        conn, pid, _ = project_db
        baseline_id = confirmed_baseline(conn, pid, "1200000", tax_basis="incl_tax")
        period_id = make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        result = cb.compare_upward_result(
            conn, pid, baseline_id, "1300000",
            settlement_tax_basis="incl_tax", period_id=period_id)
        assert result["status"] == "FAIL"
        wb = Workbook()
        wb.remove(wb.active)
        excel_export.export_control_conclusions(conn, pid, wb)
        excel_export.export_rate_rules_sheet(conn, pid, wb)
        assert "对上控制基准结论" in wb.sheetnames
        assert "费率规则与试算" in wb.sheetnames
        ws = wb["对上控制基准结论"]
        texts = {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}
        assert any("FAIL" in t for t in texts)
        assert any("不构成违规" in t for t in texts)
        assert str(result["evidence_id"]) in texts
        ws2 = wb["费率规则与试算"]
        texts2 = {str(c.value) for row in ws2.iter_rows() for c in row if c.value is not None}
        assert any("暂无费率候选" in t for t in texts2)
        assert any("不等于免计取" in t for t in texts2)

    def test_excel_conclusion_sheet_empty_project_fail_closed(self, project_db):
        conn, pid, _ = project_db
        wb = Workbook()
        wb.remove(wb.active)
        excel_export.export_control_conclusions(conn, pid, wb)
        texts = {
            str(c.value) for row in wb["对上控制基准结论"].iter_rows()
            for c in row if c.value is not None
        }
        assert any("无结论不等于通过" in t for t in texts)

    def test_markdown_export_registers_artifact(self, project_db):
        """结论报告落盘并登记为受控成果（export_runs.kind=conclusions_markdown）。"""
        conn, pid, pdir = project_db
        baseline_id = confirmed_baseline(conn, pid, "1200000")
        make_period(conn, pid, 1, "对上第1期", "upward", tax_mode="incl_tax")
        cb.compare_upward_result(
            conn, pid, baseline_id, "1100000",
            settlement_tax_basis="incl_tax")
        path = conclusions_report.export_conclusions_report(
            conn, pid, pdir / "exports")
        assert path.is_file()
        content = path.read_text(encoding="utf-8")
        assert "对上控制基准结论" in content
        rows = conn.execute(
            """SELECT kind, status FROM export_runs
               WHERE project_id=? AND kind='conclusions_markdown'""",
            (pid,)).fetchall()
        assert rows and rows[0]["status"] == "current"
