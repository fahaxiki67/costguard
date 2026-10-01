"""合成反例：单位守恒、归组行集守恒、镜像成员边界。"""
import json
from decimal import Decimal

import pytest

from jiadun.core.db import migrations
from jiadun.core.engine.aggregate import aggregate_project
from jiadun.core.matching.matching import match_items
from jiadun.core.matching.mirror import _aggregate_side, _load_match_rows

D = Decimal


@pytest.fixture()
def db(tmp_path):
    path = tmp_path / 'project.db'
    migrations.migrate(path, tmp_path / 'backups')
    conn = migrations.connect(path)
    with conn:
        pid = conn.execute("INSERT INTO projects(name,schema_version,workspace_path,created_at) VALUES ('synthetic',1,'/synthetic','2026')").lastrowid
        periods = [conn.execute("INSERT INTO settlement_periods(project_id,period_no,title,direction) VALUES (?,?,?,'downward')", (pid, n, f'第{n}期')).lastrowid for n in (1, 2, 3)]
    yield conn, pid, periods
    conn.close()


def add(conn, period, *, code='C1', name='钢材', feature='Q235', unit='吨', qty='1', price='1000', amount='1000'):
    with conn:
        return conn.execute('INSERT INTO line_items(period_id,code,name,feature,unit,quantity,unit_price,amount) VALUES (?,?,?,?,?,?,?,?)', (period,code,name,feature,unit,qty,price,amount)).lastrowid


def test_mass_conversion_and_money_conservation(db):
    conn, pid, ps = db
    add(conn, ps[0])
    add(conn, ps[1], unit='千克', qty='500', price='1', amount='500')
    agg = aggregate_project(conn, pid, direction='downward')[0]
    assert agg.cum_qty == D('1.5')
    assert agg.cum_amount == D('1500')
    assert agg.wavg_price == D('1000')
    assert agg.unit == 't'
    assert conn.execute('SELECT quantity,unit FROM line_items ORDER BY id').fetchall()[1]['quantity'] == '500'


@pytest.mark.parametrize('change', [{'unit':'m3'}, {'feature':'Q355'}, {'unit':''}])
def test_incompatible_or_missing_unit_blocks_quantity(db, change):
    conn, pid, ps = db
    add(conn, ps[0])
    add(conn, ps[1], **change)
    agg = aggregate_project(conn, pid, direction='downward')[0]
    assert agg.cum_qty is None
    assert agg.wavg_price is None
    assert agg.status != 'ok'
    assert agg.cum_amount == D('2000')


def test_fuzzy_merge_keeps_each_row_once_and_units(db):
    conn, pid, ps = db
    ids = [add(conn, ps[0], code='', name='水泥砂浆楼地面找平层施工及工程项目施工区域同一部位标准厚度工程量清单A', unit='m2'), add(conn, ps[1], code='', name='水泥砂浆楼地面找平层施工及工程项目施工区域同一部位标准厚度工程量清单B', unit='m3')]
    groups = match_items(conn, pid)
    members = [iid for g in groups for iid in g.item_ids]
    assert sorted(members) == sorted(ids)
    assert len(groups) == 1
    assert groups[0].level == 'incomparable'


def test_chained_name_merge_keeps_all_members_once(db):
    conn, pid, ps = db
    ids = [add(conn, ps[0], code='A', name='项目甲'), add(conn, ps[1], code='B', name='项目甲'), add(conn, ps[2], code='B', name='项目乙'), add(conn, ps[0], code='C', name='项目乙')]
    groups = match_items(conn, pid)
    members = [iid for g in groups for iid in g.item_ids]
    assert sorted(members) == sorted(ids)


def test_mirror_reads_only_explicit_members(db):
    conn, pid, ps = db
    a = add(conn, ps[0], code='A')
    b = add(conn, ps[1], code='B')
    add(conn, ps[2], code='A')
    rows = _load_match_rows(conn, pid, {'group_key':'downward:code:A','item_ids_json':json.dumps([a,b])})
    assert [r['id'] for r in rows] == [a,b]
    assert _load_match_rows(conn, pid, {'group_key':'code:A','item_ids_json':'[]'}) == []


def test_mirror_equal_quantities_sum_not_deduplicate(db):
    conn, pid, ps = db
    a = add(conn, ps[0], qty='60', price='1', amount='60')
    b = add(conn, ps[0], qty='60', price='1', amount='60')
    rows = _load_match_rows(conn, pid, {'group_key':'code:C1','item_ids_json':json.dumps([a,b])})
    side = _aggregate_side(rows, 'downward')
    assert side['quantity'] == D('120')
    assert side['amount'] == D('120')


def test_save_rejects_overlapping_members(db):
    from jiadun.core.matching.matching import MatchGroup, save_matches
    conn, pid, ps = db
    iid = add(conn, ps[0])
    groups = [MatchGroup(k, 'probable', 'synthetic', 0.8, item_ids=[iid]) for k in ('A','B')]
    with pytest.raises(ValueError, match='只能属于'):
        save_matches(conn, pid, groups)
    assert conn.execute('SELECT COUNT(*) FROM matches').fetchone()[0] == 0


@pytest.mark.parametrize('unit,qty,expected', [('100m²','2','200'),('10m3','3','30'),('公斤','500','0.5')])
def test_scaled_units(db, unit, qty, expected):
    conn, pid, ps = db
    add(conn, ps[0], unit=unit, qty=qty, price='1', amount=qty)
    agg = aggregate_project(conn, pid, direction='downward')[0]
    assert agg.cum_qty == D(expected)
    assert agg.wavg_price * agg.cum_qty == D(qty)


def test_difference_export_uses_complete_standard_quantities(db):
    from openpyxl import Workbook

    from jiadun.core.export.excel_export import export_diff_sheets
    conn, pid, ps = db
    add(conn, ps[0])
    add(conn, ps[0], unit='kg', qty='500', price='1', amount='500')
    add(conn, ps[1], qty='2', price='1000', amount='2000')
    add(conn, ps[2], qty=None)
    wb = Workbook()
    export_diff_sheets(conn, pid, wb)
    sheet = wb['工程量差异表']
    assert sheet.cell(2, 5).value == D('1.5')
    assert sheet.cell(3, 6).value == D('1.5')
    assert sheet.cell(4, 5).value is None
    assert sheet.cell(4, 6).value is None
    assert sheet.cell(2, 9).value == 't'


def test_difference_export_breaks_on_feature_change(db):
    from openpyxl import Workbook

    from jiadun.core.export.excel_export import export_diff_sheets
    conn, pid, ps = db
    add(conn, ps[0])
    add(conn, ps[1], feature='Q355')
    wb = Workbook()
    export_diff_sheets(conn, pid, wb)
    assert wb['工程量差异表'].cell(3, 6).value is None
    assert '不可比' in wb['工程量差异表'].cell(3, 7).value
