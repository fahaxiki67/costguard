"""真实导入入口的合成反例：各文件业务第1期不能被揉成一个期次。"""
from pathlib import Path

from openpyxl import Workbook

from jiadun.core.engine import settlement_io
from jiadun.core.models import project as pm


def test_separate_files_preserve_business_period_and_reimport_identity(tmp_path):
    info = pm.create_project('分包导入隔离', tmp_path / 'ws')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        reports = []
        for party, qty in [('A', 60), ('B', 70)]:
            src = tmp_path / f'{party}第1期.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = '第1期清单'
            ws.append(['清单编码', '清单名称', '项目特征', '单位', '工程量', '综合单价', '合价'])
            ws.append(['C1', '混凝土', 'C30', 'm3', qty, 1, qty])
            wb.save(src)
            report = settlement_io.import_settlement_file(
                conn, info.project_id, Path(info.workspace_path), src,
                direction='downward', separate_files=True,
            )
            assert report.status == 'ok', report.message
            reports.append(report)
        a, b = reports
        assert a.file_id != b.file_id
        assert a.period_id != b.period_id
        assert a.business_period_nos == {a.period_id: 1}
        assert b.business_period_nos == {b.period_id: 1}
        assert settlement_io.business_period_numbers(conn, info.project_id) == {
            a.period_id: 1, b.period_id: 1,
        }
        again = settlement_io.import_settlement_file(
            conn, info.project_id, Path(info.workspace_path), tmp_path / 'B第1期.xlsx',
            direction='downward', separate_files=True,
        )
        assert again.period_id == b.period_id
        assert again.business_period_nos == b.business_period_nos
        explicit = settlement_io.import_settlement_file(
            conn, info.project_id, Path(info.workspace_path), tmp_path / 'B第1期.xlsx',
            period_no=1, direction='downward', separate_files=True,
        )
        assert explicit.period_id == b.period_id
        assert conn.execute('SELECT COUNT(*) FROM line_items').fetchone()[0] == 2
        sources = conn.execute('SELECT source_file_id FROM settlement_periods ORDER BY id').fetchall()
        assert [row['source_file_id'] for row in sources] == [a.file_id, b.file_id]
        assert conn.execute(
            'SELECT COUNT(*) FROM line_items li JOIN settlement_periods sp ON sp.id=li.period_id '
            'JOIN raw_sheets rs ON rs.id=li.sheet_id JOIN parse_batches pb ON pb.id=rs.batch_id '
            'WHERE pb.file_id != sp.source_file_id'
        ).fetchone()[0] == 0
    finally:
        conn.close()


def test_sheet_business_periods_reuse_one_file_period(tmp_path):
    info = pm.create_project('工作表业务期号', tmp_path / 'ws')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        wb = Workbook()
        wb.remove(wb.active)
        for title, qty in [('第1期1号楼', 60), ('第1期2号楼', 70), ('第2期1号楼', 80)]:
            ws = wb.create_sheet(title)
            ws.append(['清单编码', '清单名称', '项目特征', '单位', '工程量', '综合单价', '合价'])
            ws.append(['C1', '混凝土', 'C30', 'm3', qty, 1, qty])
        src = tmp_path / 'A结算清单.xlsx'
        wb.save(src)
        report = settlement_io.import_settlement_file(
            conn, info.project_id, Path(info.workspace_path), src,
            direction='downward', separate_files=True,
        )
        assert report.status == 'ok', report.message
        assert sorted(report.business_period_nos.values()) == [1, 2]
        rows = conn.execute('SELECT period_id FROM line_items ORDER BY id').fetchall()
        assert rows[0]['period_id'] == rows[1]['period_id'] != rows[2]['period_id']
    finally:
        conn.close()


def test_manual_sheet_confirmation_keeps_source_business_period(tmp_path):
    info = pm.create_project('人工确认隔离', tmp_path / 'ws')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        period_ids = []
        for party, qty in [('A', 60), ('B', 70)]:
            src = tmp_path / f'{party}第1期.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = '汇总清单'
            ws.append(['清单编码', '清单名称', '项目特征', '单位', '工程量', '综合单价', '合价'])
            ws.append(['C1', '混凝土', 'C30', 'm3', qty, 1, qty])
            wb.save(src)
            report = settlement_io.import_settlement_file(
                conn, info.project_id, Path(info.workspace_path), src,
                direction='downward', separate_files=True,
            )
            sheet = conn.execute('SELECT id FROM raw_sheets WHERE batch_id=?', (report.batch_id,)).fetchone()
            assert settlement_io.confirm_sheet_role_and_extract(
                conn, info.project_id, sheet['id'], 'user', '合成资料确认明细范围',
                direction='downward', period_no=1,
            ) == 1
            period_ids.append(conn.execute('SELECT period_id FROM raw_sheets WHERE id=?', (sheet['id'],)).fetchone()['period_id'])
        assert len(set(period_ids)) == 2
        assert settlement_io.business_period_numbers(conn, info.project_id) == {
            period_ids[0]: 1, period_ids[1]: 1,
        }
    finally:
        conn.close()
