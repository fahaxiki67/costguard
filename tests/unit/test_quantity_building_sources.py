"""楼栋候选来源与规则自动识别（v0.1.36）反例测试。

先反例后实现：
- 仅 Sheet/文件名含楼栋（行名无）时必须有每行可追溯候选（旧实现漏）；
- 同期不同楼栋的不同行不得误判冲突；行/Sheet/文件来源矛盾必须标冲突；
- 明确且一致的单栋来源允许 rule_accepted 自动进入核量（区别于人工确认）；
- 人工确认优先覆盖自动候选；组合楼栋不摊分；
- 验收：真实合成 xlsx 导入，仅 Sheet 含 1号/2号，期次确认后无需逐行
  手填楼栋即得 1号楼 1100 超 100、全项目 1700 未超。

settings 纪律：autouse fixture 把 ``pm._SETTINGS_FILE`` 重定向到临时
settings.json（写 {}），任何 ``create_project``/导入登记不触真实 settings。
"""
import json
from decimal import Decimal

import pytest
from openpyxl import Workbook

from jiadun.core.db import migrations
from jiadun.core.engine import quantity_control as qc
from jiadun.core.engine import settlement_io


@pytest.fixture(autouse=True)
def _settings_sandbox(tmp_path, monkeypatch):
    from jiadun.core.models import project as pm

    sandbox = tmp_path / "settings-sandbox"
    sandbox.mkdir(exist_ok=True)
    settings_file = sandbox / "settings.json"
    settings_file.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(pm, "_SETTINGS_FILE", settings_file)
    yield


@pytest.fixture()
def db(tmp_path):
    path = tmp_path / "project.db"
    migrations.migrate(path, tmp_path / "backups")
    conn = migrations.connect(path)
    with conn:
        pid = conn.execute(
            "INSERT INTO projects(name,schema_version,workspace_path,created_at)"
            " VALUES (?,?,?,?)",
            ("楼栋来源", migrations.LATEST_SCHEMA_VERSION, str(tmp_path), "2026"),
        ).lastrowid
    yield conn, int(pid), tmp_path
    conn.close()


def make_sheet_meta(conn, pid, period_id, sheet_name, file_name):
    """建 raw_sheets/file 链路，返回 sheet_id（模拟导入保真层）。"""
    with conn:
        file_id = conn.execute(
            """INSERT INTO source_files(project_id, original_path, stored_path,
               original_name, sha256, size_bytes, file_type, imported_at)
               VALUES (?,?,?,?,?,?,?,?)""",
            (pid, f"/f/{file_name}", f"/s/{file_name}", file_name,
             f"{file_name}:{sheet_name}", 1, "xlsx", "2026")).lastrowid
        batch = conn.execute(
            "INSERT INTO parse_batches(file_id,parser,parsed_at,status)"
            " VALUES (?,?,?,?)", (file_id, "xlsx", "2026", "ok")).lastrowid
        sheet_id = conn.execute(
            "INSERT INTO raw_sheets(batch_id,sheet_index,sheet_name,n_rows,n_cols)"
            " VALUES (?,?,?,?,?)", (batch, 0, sheet_name, 3, 7)).lastrowid
        conn.execute("UPDATE settlement_periods SET source_file_id=? WHERE id=?",
                     (file_id, period_id))
    return sheet_id


def sheet_line(conn, period_id, sheet_id, qty, *, name="混凝土", row=None):
    flags = json.dumps({"row": row}) if row is not None else "{}"
    return conn.execute(
        "INSERT INTO line_items(period_id,sheet_id,code,name,feature,unit,"
        "quantity,flags_json) VALUES (?,?,?,?,?,?,?,?)",
        (period_id, sheet_id, "C1", name, "C30", "m3", qty, flags)).lastrowid


def confirmed_context(conn, pid, *, direction="downward", contract_key="C-A",
                      doc_kind="settlement", work_scope="材料供应"):
    context = qc.create_period_context(
        conn, pid, direction=direction, contract_key=contract_key,
        business_period_no=1, unit_name="单位", doc_kind=doc_kind,
        work_scope=work_scope, actor="t", reason="测试登记")
    qc.confirm_period_context(conn, pid, context["context_id"], actor="t",
                              reason="测试确认")
    return context


def control_of(ledger, *, code="C1", work_scope="材料供应", building=None):
    for ctl in ledger["quantity_controls"]:
        if (ctl["code"] == code and ctl["work_scope"] == work_scope
                and ctl["building"] == building):
            return ctl
    return None


# ---------------------------------------------------------------- 候选来源

class TestCandidateSources:
    def test_sheet_only_candidates_per_row(self, db):
        """仅 Sheet 名含 1号/2号（行名无）：必须有每行可追溯候选（旧实现漏）。"""
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s1 = make_sheet_meta(conn, pid, context["period_id"], "第1期1号楼", "A结算.xlsx")
        s2 = make_sheet_meta(conn, pid, context["period_id"], "第1期2号楼", "A结算.xlsx")
        sheet_line(conn, context["period_id"], s1, "100", row=2)
        sheet_line(conn, context["period_id"], s2, "200", row=2)
        candidates = qc.suggest_building_candidates(conn, pid)
        tokens = {c["building"] for c in candidates}
        assert {"1号楼", "2号楼"} <= tokens, f"got={tokens}"
        for c in candidates:
            assert {s["field"] for s in c["sources"]} <= {"name", "feature", "code", "sheet", "file"}
        one = next(c for c in candidates if c["building"] == "1号楼")
        sheet_source = next(s for s in one["sources"] if s["field"] == "sheet")
        assert sheet_source["text"] == "第1期1号楼"
        assert sheet_source["line_item_id"] > 0

    def test_same_period_different_rows_not_conflict(self, db):
        """同期两行分属不同楼栋 Sheet：各自候选，不得互相判冲突。"""
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s1 = make_sheet_meta(conn, pid, context["period_id"], "第1期1号楼", "A.xlsx")
        s2 = make_sheet_meta(conn, pid, context["period_id"], "第1期2号楼", "A.xlsx")
        sheet_line(conn, context["period_id"], s1, "100", row=2)
        sheet_line(conn, context["period_id"], s2, "200", row=2)
        candidates = qc.suggest_building_candidates(conn, pid)
        assert candidates, "应有候选"
        assert all(not c["conflict"] for c in candidates)

    def test_row_sheet_conflict_flagged(self, db):
        """行名写 1号楼、Sheet 写 2号楼：该行标冲突，不产出规则候选。"""
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s = make_sheet_meta(conn, pid, context["period_id"], "第1期2号楼", "A.xlsx")
        sheet_line(conn, context["period_id"], s, "100", name="1号楼混凝土", row=2)
        candidates = qc.suggest_building_candidates(conn, pid)
        assert candidates and all(c["conflict"] for c in candidates)

    def test_file_level_candidate(self, db):
        """仅文件名含楼栋：file 字段候选。"""
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s = make_sheet_meta(conn, pid, context["period_id"], "清单", "3号楼结算.xlsx")
        sheet_line(conn, context["period_id"], s, "100", row=2)
        candidates = qc.suggest_building_candidates(conn, pid)
        assert any(
            c["building"] == "3号楼"
            and any(src["field"] == "file" for src in c["sources"])
            for c in candidates), str(candidates)


# ---------------------------------------------------------------- 自动识别进核量

class TestRuleAccepted:
    def test_rule_accepted_enters_ledger_without_manual_building(self, db):
        """行/Sheet 来源一致且唯一 → rule_accepted 自动进单栋核量。"""
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s1 = make_sheet_meta(conn, pid, context["period_id"], "第1期1号楼", "A.xlsx")
        s2 = make_sheet_meta(conn, pid, context["period_id"], "第1期2号楼", "A.xlsx")
        sheet_line(conn, context["period_id"], s1, "100", row=2)
        sheet_line(conn, context["period_id"], s2, "200", row=2)
        ledger = qc.build_quantity_ledger(conn, pid)
        it = next(x for x in ledger["items"] if x["code"] == "C1")
        assert Decimal(it["buildings"]["1号楼"]["quantity"]) == Decimal("100")
        assert Decimal(it["buildings"]["2号楼"]["quantity"]) == Decimal("200")
        assert Decimal(it["buildings"]["__all__"]["quantity"]) == Decimal("300")
        src1 = next(s for s in it["sources"] if s["building"] == "1号楼")
        assert src1["building_status"] == "rule_accepted"
        assert "rule" in src1["building_basis"] or "自动" in src1["building_basis"]
        # 来源回溯解释字段与文本
        assert "第1期1号楼" in src1["building_basis"]
        assert not any(s["building_status"] == "human" for s in it["sources"])

    def test_row_source_disagreement_stays_pending(self, db):
        """行/Sheet 矛盾：该行楼栋 pending，不进单栋合计。"""
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s = make_sheet_meta(conn, pid, context["period_id"], "第1期2号楼", "A.xlsx")
        sheet_line(conn, context["period_id"], s, "100", name="1号楼混凝土", row=2)
        ledger = qc.build_quantity_ledger(conn, pid)
        it = next(x for x in ledger["items"] if x["code"] == "C1")
        assert "1号楼" not in it["buildings"] or it["buildings"]["1号楼"]["quantity"] is None
        assert Decimal(it["buildings"]["__all__"]["quantity"]) == Decimal("100")
        src = it["sources"][0]
        assert src["building_status"] != "rule_accepted"
        assert "冲突" in (src["building_basis"] or "") or src["building_basis"] is None

    def test_human_confirmation_overrides_auto(self, db):
        """人工确认楼栋优先于 Sheet 自动候选。"""
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s = make_sheet_meta(conn, pid, context["period_id"], "第1期1号楼", "A.xlsx")
        line_id = sheet_line(conn, context["period_id"], s, "100", row=2)
        qc.set_line_context(conn, pid, line_id, building="2号楼",
                            actor="t", reason="人工改判")
        qc.confirm_line_context(conn, pid, line_id, actor="t", reason="确认")
        ledger = qc.build_quantity_ledger(conn, pid)
        it = next(x for x in ledger["items"] if x["code"] == "C1")
        assert Decimal(it["buildings"]["2号楼"]["quantity"]) == Decimal("100")
        assert it["buildings"].get("1号楼", {}).get("quantity") in (None, "0") or \
            "1号楼" not in it["buildings"]
        src = it["sources"][0]
        assert src["building_status"] == "line_confirmed"

    def test_combined_scope_not_split_or_averaged(self, db):
        """Sheet 为组合范围 1-3号楼：全项目计入，单栋 1号楼不受影响。"""
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s = make_sheet_meta(conn, pid, context["period_id"], "第1期1-3号楼", "A.xlsx")
        sheet_line(conn, context["period_id"], s, "300", row=2)
        ledger = qc.build_quantity_ledger(conn, pid)
        it = next(x for x in ledger["items"] if x["code"] == "C1")
        assert Decimal(it["buildings"]["__all__"]["quantity"]) == Decimal("300")
        assert "1号楼" not in it["buildings"] or it["buildings"]["1号楼"]["quantity"] is None
        assert it["row_scope"]["multi_building_scopes"] == ["1-3号楼"]


# ---------------------------------------------------------------- 导入验收

class TestImportAcceptance:
    def _import_contract(self, conn, pid, ws, key, direction, kind, category, rows):
        """同 file 两 Sheet（仅 Sheet 名含楼栋），U1000/800 A600/200 B500/300 C0/100。"""
        source = ws / f"{key}第1期.xlsx"
        wb = Workbook()
        wb.remove(wb.active)
        for building, qty in rows:
            sheet = wb.create_sheet(f"第1期{building}")
            sheet.append(["清单编码", "清单名称", "项目特征", "单位", "工程量",
                          "综合单价", "合价"])
            sheet.append(["C1", "混凝土", "C30", "m3", qty, 1, qty])
        wb.save(source)
        report = settlement_io.import_settlement_file(
            conn, pid, ws, source, direction=direction,
            separate_files=True, document_category=category)
        assert report.status == "ok", report.message
        context_id = qc.attach_period_context(
            conn, pid, report.period_id, contract_key=key, business_period_no=1,
            unit_name=f"单位{key}", doc_kind=kind, work_scope="材料供应",
            actor="user", reason="期次身份确认")
        qc.confirm_period_context(conn, pid, context_id, actor="user",
                                  reason="期次身份确认")
        return report

    def test_building_auto_without_per_row_confirmation(self, db):
        conn, pid, ws = db
        self._import_contract(conn, pid, ws, "U", "upward", "contract",
                              "upward_contract_boq",
                              [("1号楼", 1000), ("2号楼", 800)])
        for key, rows in (("A", [("1号楼", 600), ("2号楼", 200)]),
                          ("B", [("1号楼", 500), ("2号楼", 300)]),
                          ("C", [("1号楼", 0), ("2号楼", 100)])):
            self._import_contract(conn, pid, ws, key, "downward", "settlement",
                                  "downward_settlement", rows)
        ledger = qc.build_quantity_ledger(conn, pid)
        b1 = control_of(ledger, building="1号楼")
        assert Decimal(b1["downstream_quantity"]) == Decimal("1100"), b1["downstream_quantity"]
        assert Decimal(b1["delta_vs_contract"]) == Decimal("100")
        assert b1["status_vs_contract"] == "FAIL"
        whole = control_of(ledger, building=None)
        assert Decimal(whole["downstream_quantity"]) == Decimal("1700")
        assert Decimal(whole["baselines"]["upward_contract"]["quantity"]) == Decimal("1800")
        assert whole["status_vs_contract"] == "PASS"
        # 未逐行手填：来源应标记规则自动而非人工
        it = next(x for x in ledger["items"]
                  if x["code"] == "C1" and x["doc_kind"] == "settlement")
        assert all(s["building_status"] == "rule_accepted" for s in it["sources"])
        # 楼栋候选按行区分
        candidates = qc.suggest_building_candidates(conn, pid)
        by_building = {}
        for c in candidates:
            by_building.setdefault(c["building"], set()).update(
                s["line_item_id"] for s in c["sources"])
        assert by_building["1号楼"].isdisjoint(by_building["2号楼"])

    def test_combined_sheet_in_import_not_split(self, db):
        conn, pid, ws = db
        source = ws / "合计第1期.xlsx"
        wb = Workbook()
        wb.remove(wb.active)
        sheet = wb.create_sheet("第1期1-3号楼")
        sheet.append(["清单编码", "清单名称", "项目特征", "单位", "工程量",
                      "综合单价", "合价"])
        sheet.append(["C1", "混凝土", "C30", "m3", 300, 1, 300])
        wb.save(source)
        report = settlement_io.import_settlement_file(
            conn, pid, ws, source, direction="downward", separate_files=True,
            document_category="downward_settlement")
        assert report.status == "ok", report.message
        context_id = qc.attach_period_context(
            conn, pid, report.period_id, contract_key="D", business_period_no=1,
            unit_name="单位D", doc_kind="settlement", work_scope="材料供应",
            actor="user", reason="确认")
        qc.confirm_period_context(conn, pid, context_id, actor="user",
                                  reason="确认")
        ledger = qc.build_quantity_ledger(conn, pid)
        it = next(x for x in ledger["items"] if x["code"] == "C1")
        assert Decimal(it["buildings"]["__all__"]["quantity"]) == Decimal("300")
        assert "1号楼" not in it["buildings"] or it["buildings"]["1号楼"]["quantity"] is None
        assert it["row_scope"]["multi_building_scopes"] == ["1-3号楼"]


class TestTokenBoundaries:
    """P0 安全边界：中文序数/区间/并列不得误识（rule_accepted 前提）。"""

    @pytest.mark.parametrize("text,expected", [
        ("十一号楼", ["11号楼"]),
        ("二十号楼", ["20号楼"]),
        ("二十一号楼", ["21号楼"]),
        ("一至三号楼", ["1-3号楼"]),
        ("1、2号楼", ["1、2号楼"]),
        ("一百零一号楼", []),
        ("101号楼", ["101号楼"]),
        ("1号楼", ["1号楼"]),
        ("1#楼", ["1号楼"]),
        ("一号楼", ["1号楼"]),
        ("1-3号楼", ["1-3号楼"]),
    ])
    def test_token_extraction(self, text, expected):
        assert qc._building_tokens(text) == expected

    def test_cn_ordinal_building_rule_accepted(self, db):
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s = make_sheet_meta(conn, pid, context["period_id"], "十一号楼清单", "A.xlsx")
        sheet_line(conn, context["period_id"], s, "100", row=2)
        ledger = qc.build_quantity_ledger(conn, pid)
        it = next(x for x in ledger["items"] if x["code"] == "C1")
        assert Decimal(it["buildings"]["11号楼"]["quantity"]) == Decimal("100")
        src = it["sources"][0]
        assert src["building_status"] == "rule_accepted"
        assert "11号楼" in src["building_basis"]

    def test_cn_range_and_parallel_are_combined_only(self, db):
        conn, pid, _ = db
        context = confirmed_context(conn, pid)
        s1 = make_sheet_meta(conn, pid, context["period_id"], "一至三号楼汇总", "A.xlsx")
        s2 = make_sheet_meta(conn, pid, context["period_id"], "1、2号楼清单", "B.xlsx")
        sheet_line(conn, context["period_id"], s1, "100", row=2)
        sheet_line(conn, context["period_id"], s2, "50", row=2)
        ledger = qc.build_quantity_ledger(conn, pid)
        it = next(x for x in ledger["items"] if x["code"] == "C1")
        assert Decimal(it["buildings"]["__all__"]["quantity"]) == Decimal("150")
        # 组合范围不产生/不影响任何单栋结论
        assert "1号楼" not in it["buildings"] or \
            it["buildings"]["1号楼"]["quantity"] is None
        assert set(it["row_scope"]["multi_building_scopes"]) == {"1-3号楼", "1、2号楼"}
        for src in it["sources"]:
            assert src["building_status"] == "combined_scope"


@pytest.mark.parametrize("text,expected", [
    ("十一一号楼", []),
    ("1、2号楼及3、4号楼及5号楼", ["1、2号楼", "3、4号楼", "5号楼"]),
])
def test_additional_token_boundaries(text, expected):
    assert qc._building_tokens(text) == expected


def test_distinct_combined_sources_are_conflict():
    fields = qc._row_building_source_fields(None, None, None, "1-3号楼", "4-6号楼.xlsx")
    building, status, basis = qc._derive_rule_building(fields)
    assert building is None and status is None
    assert "冲突" in basis


def test_combined_and_single_candidate_is_conflict(db):
    conn, pid, _ = db
    context = confirmed_context(conn, pid)
    sheet = make_sheet_meta(conn, pid, context['period_id'], '1-3号楼', 'A.xlsx')
    line_id = sheet_line(conn, context['period_id'], sheet, '100', name='1号楼混凝土')
    candidates = qc.suggest_building_candidates(conn, pid)
    assert candidates and all(c['conflict'] for c in candidates if c['line_item_id'] == line_id)
    source = qc.build_quantity_ledger(conn, pid)['items'][0]['sources'][0]
    assert source['building'] is None
    assert '并存' in source['building_basis']
