"""主代理独立复算留下的实际口径反例。"""
from pathlib import Path

from jiadun.core.engine import quantity_control as qc
from jiadun.core.models import project as pm


def test_pending_period_line_scope_override_blocks_matching_scope(tmp_path):
    info = pm.create_project('行口径覆盖反例', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        for direction, key, kind, scope, quantity, confirmed in [
            ('upward', 'U', 'contract', '材料供应', '100', True),
            ('downward', 'A', 'settlement', '材料供应', '60', True),
            ('downward', 'B', 'settlement', '劳务', '60', False),
        ]:
            context = qc.create_period_context(
                conn, info.project_id, direction=direction, contract_key=key,
                business_period_no=1, unit_name=key, doc_kind=kind,
                work_scope=scope, actor='user', reason='合成范围登记',
            )
            line_id = conn.execute(
                'INSERT INTO line_items(period_id,code,name,feature,unit,quantity,flags_json) VALUES (?,?,?,?,?,?,?)',
                (context['period_id'], 'C1', '混凝土', 'C30', 'm3', quantity, '{"row":2}'),
            ).lastrowid
            if key == 'B':
                qc.set_line_context(conn, info.project_id, line_id, work_scope='材料供应',
                                    actor='user', reason='行计量口径已明确，期次尚待确认')
                qc.confirm_line_context(conn, info.project_id, line_id, actor='user', reason='确认该行材料口径')
            if confirmed:
                qc.confirm_period_context(conn, info.project_id, context['context_id'],
                                          actor='user', reason='合成合同范围确认')
        controls = qc.build_quantity_ledger(conn, info.project_id)['quantity_controls']
        control = next(row for row in controls if row['building'] is None and row['work_scope'] == '材料供应')
        assert control['status_vs_contract'] == control['status'] == 'PENDING'
        assert len(control['excluded_details']) == 1
        assert control['delta_vs_contract'] is None
        assert control['excluded_details'][0]['line_item_id'] == line_id
        from openpyxl import Workbook

        from jiadun.core.export.excel_export import export_quantity_control_sheets
        workbook = Workbook()
        export_quantity_control_sheets(qc.build_quantity_ledger(conn, info.project_id), workbook)
        data = list(workbook['工程量控制台账'].iter_rows(values_only=True))
        exported = dict(zip(data[0], data[1], strict=True))
        assert exported['对下已结算合计'] is None
        assert exported['较合同差量'] is None
    finally:
        conn.close()


def test_unregistered_periods_remain_visible_in_quantity_exports(tmp_path):
    info = pm.create_project('旧期次待登记', tmp_path / 'workspace')
    info, conn = pm.open_project(Path(info.workspace_path))
    try:
        period_id = conn.execute(
            "INSERT INTO settlement_periods(project_id,period_no,title,direction) VALUES (?,?,?,?)",
            (info.project_id, 3, '旧结算第3期', 'downward'),
        ).lastrowid
        conn.execute("INSERT INTO line_items(period_id,name,unit,quantity) VALUES (?,?,?,?)",
                     (period_id, '混凝土', 'm3', '100'))
        ledger = qc.build_quantity_ledger(conn, info.project_id)
        assert ledger['quantity_controls'] == []
        assert ledger['unregistered_periods'][0]['period_id'] == period_id
        assert ledger['unregistered_periods'][0]['detail_count'] == 1
        from openpyxl import Workbook

        from jiadun.core.export.excel_export import export_quantity_control_sheets
        workbook = Workbook()
        export_quantity_control_sheets(ledger, workbook)
        assert workbook['核量待确认范围'].cell(2, 1).value == 'unregistered_periods'
    finally:
        conn.close()
